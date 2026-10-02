import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    import importlib
    from app import db
    importlib.reload(db)
    from app import export, importer, main
    importlib.reload(importer); importlib.reload(export); importlib.reload(main)
    with TestClient(main.app) as c:
        yield c


def post(c, user, entries=(), team=(), projects=()):
    r = c.post("/api/changes", json={"user": user, "entries": list(entries), "team": list(team), "projects": list(projects)})
    assert r.status_code == 200, r.text
    return r.json()


def test_seed(client):
    s = client.get("/api/state").json()
    assert len(s["team"]) == 5 and len(s["projects"]) == 7 and s["entries"] == []


def test_entry_lifecycle_and_log(client):
    e = {"date": "2026-10-01", "person": "KOUTSELIS Dimitris", "project": "Wiseman"}
    s = post(client, "Nikos", [{**e, "oldH": 0, "newH": 4.5}])
    assert s["entries"][0]["hours"] == 4.5 and s["conflicts"] == 0
    post(client, "Nikos", [{**e, "oldH": 4.5, "newH": 3}])
    s = post(client, "Nikos", [{**e, "oldH": 3, "newH": 0}])
    assert s["entries"] == []
    actions = [l["action"] for l in client.get("/api/log").json()]
    assert actions[:3] == ["Διαγραφή", "Αλλαγή", "Προσθήκη"]


def test_conflict_is_logged(client):
    e = {"date": "2026-10-01", "person": "Katsifi Nelly", "project": "Promitheas"}
    post(client, "A", [{**e, "oldH": 0, "newH": 2}])
    s = post(client, "B", [{**e, "oldH": 0, "newH": 5}])  # ο B δεν είχε δει την αλλαγή του A
    assert s["conflicts"] == 1 and s["entries"][0]["hours"] == 5
    assert "A" in client.get("/api/log").json()[0]["note"]


def test_validation(client):
    r = client.post("/api/changes", json={"user": "x", "entries": [
        {"date": "2026-10-01", "person": "P", "project": "J", "oldH": 0, "newH": 30}]})
    assert r.status_code == 422


def test_new_person_project_and_export(client):
    post(client, "Nikos", team=[{"name": "NEW Person", "grade": 9, "type": "DEV", "active": True}],
         projects=[{"name": "New Proj", "active": True}])
    post(client, "Nikos", [{"date": "2026-10-02", "person": "NEW Person", "project": "New Proj", "oldH": 0, "newH": 6}])
    r = client.get("/api/export?from=2026-10-01&to=2026-10-31")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "NEW Person" in wb.sheetnames and "Summary" in wb.sheetnames and "Log" in wb.sheetnames
    ws = wb["NEW Person"]
    headers = [c.value for c in ws[3]]
    assert "New Proj" in headers
    col = headers.index("New Proj") + 1
    row = next(r for r in range(4, ws.max_row) if ws.cell(r, 1).value and ws.cell(r, 1).value.day == 2)
    assert ws.cell(row, col).value == 6
