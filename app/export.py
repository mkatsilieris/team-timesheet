"""Παραγωγή Excel για download: ένα sheet ανά άτομο (ημερομηνίες × projects), σύνοψη, πρωτογενή δεδομένα και ιστορικό."""
from __future__ import annotations

import io
import re
from collections import defaultdict
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as col

from . import db

DAY_HOURS = 8
GR_DAYS = ["Δευ", "Τρι", "Τετ", "Πεμ", "Παρ", "Σαβ", "Κυρ"]
F = "Arial"
HDR_FONT = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="0F4C5C")
TOT_FILL = PatternFill("solid", fgColor="E6F0F2")
WARN_FILL = PatternFill("solid", fgColor="FCD5B4")
WK_FILL = PatternFill("solid", fgColor="F0F0F0")
BASE = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
NUM = "0.0;-0.0;"


def _sheet_title(name: str, used: set[str]) -> str:
    t = re.sub(r"[\[\]\*\?/\\:]", "", name)[:31] or "Sheet"
    base, i = t, 2
    while t.lower() in used:
        suffix = f" ({i})"
        t = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(t.lower())
    return t


def _header(ws, row: int, values: list, widths: list[int] | None = None):
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for j, w in enumerate(widths, 1):
            ws.column_dimensions[col(j)].width = w


def _table(ws, header: list, rows: list[list], widths: list[int], date_cols=(), num_cols=()):
    _header(ws, 1, header, widths)
    for i, r in enumerate(rows, 2):
        for j, v in enumerate(r, 1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = BASE
            if j - 1 in date_cols and isinstance(v, date):
                c.number_format = "dd/mm/yyyy"
            if j - 1 in num_cols:
                c.number_format = "0.00"
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{col(len(header))}{len(rows) + 1}"


def _d(s: str | None):
    return date.fromisoformat(s) if s else None


def build_export(date_from: str | None = None, date_to: str | None = None) -> tuple[bytes, str, str]:
    with db.read_conn() as conn:
        team = [dict(r) for r in conn.execute("SELECT * FROM team ORDER BY name COLLATE NOCASE")]
        projects = [dict(r) for r in conn.execute("SELECT * FROM projects ORDER BY sort, name")]
        bounds = conn.execute("SELECT MIN(date), MAX(date) FROM entries").fetchone()
        today = date.today()
        a = date_from or bounds[0] or today.replace(day=1).isoformat()
        b = date_to or bounds[1] or today.isoformat()
        entries = [dict(r) for r in conn.execute(
            "SELECT * FROM entries WHERE date BETWEEN ? AND ? ORDER BY date, person, project", (a, b))]
        log = [dict(r) for r in conn.execute("SELECT * FROM log ORDER BY id")]

    hours = defaultdict(float)  # (person, date, project) -> hours
    for e in entries:
        hours[(e["person"], e["date"], e["project"])] = e["hours"]
    proj_names = [p["name"] for p in projects]
    extra = sorted({e["project"] for e in entries} - set(proj_names))
    proj_names += extra  # projects που σβήστηκαν αλλά έχουν ώρες

    wb = Workbook()
    used: set[str] = set()

    # --- Σύνοψη ---
    ws = wb.active
    ws.title = _sheet_title("Summary", used)
    ws["A1"] = f"Ώρες ανά άτομο και project, {_d(a):%d/%m/%Y} έως {_d(b):%d/%m/%Y}"
    ws["A1"].font = Font(name=F, size=13, bold=True, color="0F4C5C")
    tot = defaultdict(float)
    for e in entries:
        tot[(e["person"], e["project"])] += e["hours"]
    used_projects = [p for p in proj_names if any(tot[(t["name"], p)] for t in team)] or proj_names
    people = [t for t in team if t["active"] or any(tot[(t["name"], p)] for p in proj_names)]
    heads = ["Άτομο", "Grade", "Type", "Σύνολο ωρών", "Man-days"] + used_projects
    for j, h in enumerate(heads, 1):
        c = ws.cell(row=3, column=j, value=h)
        c.font, c.fill, c.border = HDR_FONT, HDR_FILL, BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[3].height = 40
    first, last = 4, 3 + len(people)
    pc1 = 6
    for i, t in enumerate(people, first):
        ws.cell(row=i, column=1, value=t["name"]).font = BOLD
        ws.cell(row=i, column=2, value=t["grade"]).font = BASE
        ws.cell(row=i, column=3, value=t["type"]).font = BASE
        for j, p in enumerate(used_projects, pc1):
            ws.cell(row=i, column=j, value=tot[(t["name"], p)] or None)
        ws.cell(row=i, column=4, value=f"=SUM({col(pc1)}{i}:{col(pc1 + len(used_projects) - 1)}{i})")
        ws.cell(row=i, column=5, value=f"=D{i}/{DAY_HOURS}")
    tr = last + 1
    ws.cell(row=tr, column=1, value="Σύνολο ομάδας")
    for j in range(4, pc1 + len(used_projects)):
        ws.cell(row=tr, column=j, value=f"=SUM({col(j)}{first}:{col(j)}{last})" if people else 0)
    for r in range(first, tr + 1):
        for j in range(1, len(heads) + 1):
            c = ws.cell(row=r, column=j)
            c.border = BOX
            if c.font != BOLD:
                c.font = BASE
            if j >= 4:
                c.number_format = "#,##0.0;-#,##0.0;-"
                c.alignment = Alignment(horizontal="center")
            if r == tr:
                c.font, c.fill = BOLD, TOT_FILL
    ws.column_dimensions["A"].width = 26
    for j, w in zip(range(2, 6), (7, 7, 12, 10)):
        ws.column_dimensions[col(j)].width = w
    for j in range(pc1, pc1 + len(used_projects)):
        ws.column_dimensions[col(j)].width = 14
    ws.freeze_panes = ws.cell(row=first, column=2)

    # --- Ένα sheet ανά άτομο ---
    da, dbb = _d(a), _d(b)
    for t in people:
        p_ws = wb.create_sheet(_sheet_title(t["name"], used))
        p_ws["A1"], p_ws["B1"] = "Άτομο", t["name"]
        p_ws["D1"], p_ws["E1"] = "Grade", t["grade"]
        p_ws["F1"], p_ws["G1"] = "Type", t["type"]
        for ref in ("A1", "D1", "F1"):
            p_ws[ref].font = BOLD
        p_ws["B1"].font = Font(name=F, size=12, bold=True, color="0F4C5C")
        active = {pp["name"] for pp in projects if pp["active"]}
        with_hours = {pj for (pp, _, pj), v in hours.items() if pp == t["name"] and v}
        cols = [p for p in proj_names if p in active or p in with_hours]
        _header(p_ws, 3, ["Ημερομηνία", "Ημέρα", "Σύνολο"] + cols, [12, 7, 9] + [14] * len(cols))
        p_ws.row_dimensions[3].height = 40
        r = 4
        d = da
        days_with_hours = {dd for (pp, dd, _) in hours if pp == t["name"]}
        while d <= dbb:
            iso = d.isoformat()
            if d.weekday() < 5 or iso in days_with_hours:
                p_ws.cell(row=r, column=1, value=d).number_format = "dd/mm/yyyy"
                p_ws.cell(row=r, column=2, value=GR_DAYS[d.weekday()])
                for j, p in enumerate(cols, 4):
                    v = hours.get((t["name"], iso, p))
                    c = p_ws.cell(row=r, column=j, value=v or None)
                    c.number_format = NUM
                    c.alignment = Alignment(horizontal="center")
                total = sum(hours.get((t["name"], iso, p), 0) for p in cols)
                c = p_ws.cell(row=r, column=3, value=f"=SUM(D{r}:{col(3 + len(cols))}{r})" if cols else 0)
                c.number_format, c.font, c.alignment = NUM, BOLD, Alignment(horizontal="center")
                if total > DAY_HOURS:
                    c.fill = WARN_FILL
                for j in range(1, 4 + len(cols)):
                    cc = p_ws.cell(row=r, column=j)
                    cc.border = BOX
                    if cc.font != BOLD:
                        cc.font = BASE
                    if d.weekday() >= 5 and j != 3:
                        cc.fill = WK_FILL
                r += 1
            d += timedelta(days=1)
        last_row = r - 1
        p_ws.cell(row=r, column=1, value="Σύνολο")
        for j in range(3, 4 + len(cols)):
            p_ws.cell(row=r, column=j, value=f"=SUM({col(j)}4:{col(j)}{last_row})" if last_row >= 4 else 0)
        for j in range(1, 4 + len(cols)):
            c = p_ws.cell(row=r, column=j)
            c.font, c.fill, c.border, c.number_format = BOLD, TOT_FILL, BOX, "#,##0.0;-#,##0.0;-"
            c.alignment = Alignment(horizontal="center")
        p_ws.freeze_panes = "D4"

    # --- Πρωτογενή δεδομένα ---
    _table(wb.create_sheet(_sheet_title("Entries", used)), ["Date", "Person", "Project", "Hours", "UpdatedBy", "UpdatedAt"],
           [[_d(e["date"]), e["person"], e["project"], e["hours"], e["updated_by"], e["updated_at"]] for e in entries],
           [12, 26, 30, 8, 22, 20], date_cols=(0,), num_cols=(3,))
    monthly = defaultdict(float)
    for e in entries:
        monthly[(e["date"][:7], e["person"], e["project"])] += e["hours"]
    _table(wb.create_sheet(_sheet_title("MonthlyTotals", used)), ["Month", "Person", "Project", "Hours", "ManDays"],
           [[m, p, j, h, round(h / DAY_HOURS, 2)] for (m, p, j), h in sorted(monthly.items())],
           [10, 26, 30, 8, 9], num_cols=(3, 4))
    _table(wb.create_sheet(_sheet_title("Team", used)), ["Person", "Grade", "Type", "Active"],
           [[t["name"], t["grade"], t["type"], "Ναι" if t["active"] else "Όχι"] for t in team], [26, 8, 8, 8])
    _table(wb.create_sheet(_sheet_title("Projects", used)), ["Project", "Active"],
           [[p["name"], "Ναι" if p["active"] else "Όχι"] for p in projects], [34, 8])
    _table(wb.create_sheet(_sheet_title("Log", used)),
           ["Timestamp", "User", "Action", "Date", "Person", "Project", "OldHours", "NewHours", "Note"],
           [[l["ts"], l["user"], l["action"], _d(l["date"]), l["person"], l["project"], l["old_hours"], l["new_hours"], l["note"]]
            for l in log], [20, 22, 20, 12, 26, 30, 9, 9, 60], date_cols=(3,))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), a, b
