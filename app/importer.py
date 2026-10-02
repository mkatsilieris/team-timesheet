"""Εισαγωγή του Timesheet_DB.xlsx της προηγούμενης έκδοσης (HTML + Excel στο OneDrive)."""
from __future__ import annotations

from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from . import db

NO = {"όχι", "οχι", "no", "false", "0", "n"}


def _iso(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)):
        return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(s[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return None


def _rows(wb, name):
    if name not in wb.sheetnames:
        return []
    it = wb[name].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it, [])]
    return [dict(zip(header, r)) for r in it]


def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def import_legacy_xlsx(path: Path) -> dict:
    wb = load_workbook(path, data_only=True, read_only=True)
    counts = {"team": 0, "projects": 0, "entries": 0, "log": 0}
    with db.write_tx() as conn:
        for r in _rows(wb, "Team"):
            name = str(r.get("Person") or "").strip()
            if not name:
                continue
            grade = _num(r.get("Grade"))
            active = 0 if str(r.get("Active") or "").strip().lower() in NO else 1
            conn.execute(
                "INSERT OR REPLACE INTO team(name,grade,type,active) VALUES (?,?,?,?)",
                (name, int(grade) if grade is not None else None, str(r.get("Type") or "").strip(), active),
            )
            counts["team"] += 1
        for i, r in enumerate(_rows(wb, "Projects")):
            name = str(r.get("Project") or "").strip()
            if not name:
                continue
            active = 0 if str(r.get("Active") or "").strip().lower() in NO else 1
            conn.execute("INSERT OR REPLACE INTO projects(name,active,sort) VALUES (?,?,?)", (name, active, i))
            counts["projects"] += 1
        for r in _rows(wb, "Entries"):
            d, person, project = _iso(r.get("Date")), str(r.get("Person") or "").strip(), str(r.get("Project") or "").strip()
            hours = _num(r.get("Hours"))
            if not d or not person or not project or not hours or hours <= 0:
                continue
            conn.execute(
                "INSERT OR REPLACE INTO entries(date,person,project,hours,updated_by,updated_at) VALUES (?,?,?,?,?,?)",
                (d, person, project, hours, str(r.get("UpdatedBy") or ""), str(r.get("UpdatedAt") or "")),
            )
            counts["entries"] += 1
        for r in _rows(wb, "Log"):
            if not r.get("Timestamp"):
                continue
            db.add_log(
                conn,
                str(r.get("User") or ""),
                str(r.get("Action") or ""),
                _iso(r.get("Date")),
                str(r.get("Person") or "") or None,
                str(r.get("Project") or "") or None,
                _num(r.get("OldHours")),
                _num(r.get("NewHours")),
                str(r.get("Note") or ""),
                ts=str(r.get("Timestamp")),
            )
            counts["log"] += 1
        db.add_log(conn, "system", "Εισαγωγή από Excel", note=f"{path.name}: {counts}")
    wb.close()
    return counts
